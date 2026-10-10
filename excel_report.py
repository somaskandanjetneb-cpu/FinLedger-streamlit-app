from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment
def excel_bytes(rows,company,start,end):
    wb=Workbook(); ws=wb.active; ws.title="Transactions"
    ws.append([company["company_name"]]); ws.append([company["address"] or ""]); ws.append([f"Period: {start} to {end}"]); ws.append([])
    ws.append(["Date","Type","Category","Amount","Details","Reference","User"])
    for c in ws[1]: c.font=Font(bold=True,size=16)
    for c in ws[5]: c.font=Font(bold=True)
    for r in rows: ws.append([r["trans_date"],r["trans_type"],r["category"],float(r["amount"]),r["details"] or "",r["reference_no"] or "",r["username"] or ""])
    for row in ws.iter_rows(min_row=6): row[3].number_format='₹ #,##0.00'
    widths=[15,14,22,18,35,20,18]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    out=BytesIO(); wb.save(out); return out.getvalue()
